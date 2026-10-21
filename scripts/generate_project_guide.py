from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "多智能体工单系统-项目技术说明书.docx"
ARCHITECTURE = ROOT / "docs" / "images" / "architecture.png"
WORKFLOW = ROOT / "docs" / "images" / "workflow.png"
DESKTOP = ROOT / "docs" / "screenshots" / "copilot-dashboard-desktop.png"
CHAT = ROOT / "docs" / "screenshots" / "copilot-chat-desktop.png"
MOBILE = ROOT / "docs" / "screenshots" / "copilot-dashboard-mobile.png"

NAVY = "17324D"
BLUE = "2463EB"
LIGHT_BLUE = "EAF2FF"
LIGHT_GRAY = "F3F6F9"
LIGHT_GREEN = "E9F8F2"
GREEN = "0F7A5A"
LIGHT_AMBER = "FFF6DF"
AMBER = "A45F00"
RED = "B4232A"
TEXT = RGBColor(39, 52, 67)
MUTED = RGBColor(100, 116, 139)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(cell, text: str, bold: bool = False, color: RGBColor | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9.5)
    if color:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_toc(document: Document) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "在 Word 中右键此处并选择“更新域”生成目录"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])

    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)
        paragraph.paragraph_format.space_after = Pt(3)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        set_cell_text(header.cells[index], text, bold=True, color=RGBColor(255, 255, 255))
        set_cell_shading(header.cells[index], NAVY)
        if widths:
            header.cells[index].width = Cm(widths[index])
    for row_data in rows:
        row = table.add_row()
        for index, text in enumerate(row_data):
            set_cell_text(row.cells[index], text)
            if widths:
                row.cells[index].width = Cm(widths[index])
    document.add_paragraph()
    return table


def add_callout(document: Document, title: str, text: str, fill: str = LIGHT_BLUE) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    title_run = p.add_run(title)
    title_run.bold = True
    title_run.font.name = "Microsoft YaHei"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_run.font.size = Pt(10.5)
    body = cell.add_paragraph(text)
    body.paragraph_format.space_after = Pt(0)
    body.paragraph_format.line_spacing = 1.35
    document.add_paragraph()


def add_code(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F6F8FA")
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)
    paragraph.paragraph_format.line_spacing = 1.2
    document.add_paragraph()


def add_image(document: Document, path: Path, width_cm: float, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.style = document.styles["Caption"]


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color in (
        ("Title", 30, NAVY),
        ("Subtitle", 14, "64748B"),
        ("Heading 1", 20, NAVY),
        ("Heading 2", 14, BLUE),
        ("Heading 3", 11.5, GREEN),
    ):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = style_name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)

    caption = styles["Caption"]
    caption.font.name = "Microsoft YaHei"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(9)
    caption.font.color.rgb = MUTED

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("多智能体工单协同系统 | 项目技术说明书")
    header_run.font.name = "Microsoft YaHei"
    header_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer_run = footer.add_run("项目学习与面试复盘文档  |  ")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = MUTED
    add_page_number(footer)


def build_document() -> None:
    document = Document()
    configure_document(document)
    properties = document.core_properties
    properties.title = "多智能体工单协同系统 - 项目技术说明书"
    properties.subject = "项目架构、业务流程、测试与面试讲解"
    properties.keywords = "FastAPI, SQLite, LangGraph, CopilotKit, multi-agent, ticket system"
    properties.comments = "用于项目复盘、GitHub 说明和实习面试准备。"

    cover_spacer = document.add_paragraph()
    cover_spacer.paragraph_format.space_after = Pt(60)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("多智能体工单协同系统")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("项目组成、工作原理与面试讲解手册")

    add_callout(
        document,
        "文档用途",
        "帮助项目作者系统理解代码结构、端到端流程、关键技术决策、测试方式和设计边界，"
        "并能够在实习面试中准确讲解项目，不夸大实现范围。",
        LIGHT_BLUE,
    )
    add_table(
        document,
        ["项目信息", "内容"],
        [
            ["项目名称", "多智能体工单协同系统"],
            ["主要技术", "FastAPI、SQLite、LangGraph、CopilotKit、LlamaIndex、BM25"],
            ["项目形态", "本地可运行的单端口 Web 应用"],
            ["文档版本", f"1.0 / {date.today().isoformat()}"],
            ["适用场景", "GitHub 项目说明、实习面试复盘、代码答辩准备"],
        ],
        widths=[4.0, 11.5],
    )
    document.add_page_break()

    document.add_heading("目录", level=1)
    add_toc(document)
    document.add_page_break()

    document.add_heading("1. 项目概览", level=1)
    document.add_heading("1.1 项目定位", level=2)
    document.add_paragraph(
        "这是一个面向企业内部多部门协作的工单平台。用户可以用一段自然语言同时提出多个需求，"
        "系统将父工单拆分成多个部门子工单，交给 IT、HR、采购、财务或通用 Agent 处理，"
        "最后汇总结果并实时回显到页面。"
    )
    document.add_paragraph(
        "项目不是单纯的表单 CRUD，也不是只展示模型对话的演示页。它包含完整的接入层、状态模型、"
        "任务分派、缺失信息补充、SLA 升级、知识检索、人工确认和浏览器回归测试。"
    )
    document.add_heading("1.2 解决的问题", level=2)
    add_bullets(
        document,
        [
            "跨部门工单靠人工转发，容易遗漏、重复沟通、无法追踪。",
            "用户一句话可能包含多个独立诉求，传统表单只能提交到一个部门。",
            "硬件、采购、报销等工单缺少关键字段时，处理人需要反复询问。",
            "没有实时状态和 SLA 机制时，工单容易长期停留在处理中。",
            "页面助手如果可以直接提交，会带来误操作和数据写入风险。",
        ],
    )
    document.add_heading("1.3 本次实现范围", level=2)
    add_table(
        document,
        ["已实现", "明确边界"],
        [
            ["工单创建、查询、编辑、删除", "数据仍使用 SQLite，没有迁移 MySQL"],
            ["多意图分类、父子工单拆分", "模型调用失败时使用关键词规则"],
            ["部门 Agent 与结果聚合", "Agent 使用进程内线程，未接入消息队列"],
            ["等待补充信息、SLA 升级", "IT、HR、库存等业务数据为本地演示数据"],
            ["LangGraph 草稿分析与知识检索", "技能目录主要提供元数据，核心规则由代码显式实现"],
            ["CopilotKit 表单回填与人工确认", "生产环境需替换演示 Token 接口"],
        ],
        widths=[7.5, 8.0],
    )

    document.add_heading("2. 快速认识项目", level=1)
    document.add_heading("2.1 核心能力", level=2)
    add_table(
        document,
        ["能力", "输入", "输出"],
        [
            ["多意图识别", "一段自由文本", "IT/HR/采购/财务等多个意图"],
            ["父工单拆分", "分类结果", "一个父工单和多个部门子工单"],
            ["部门处理", "子工单与元数据", "结构化处理结果或补充信息请求"],
            ["SLA 监控", "子工单创建时间和部门阈值", "升级状态与聚合提示"],
            ["草稿分析", "当前表单内容", "可编辑草稿、缺失字段、SOP 建议"],
            ["页面助手", "对话消息和表单快照", "fillTicketForm 工具调用"],
            ["实时展示", "工单状态变化", "WebSocket 推送与前端刷新"],
        ],
        widths=[4.0, 5.5, 6.0],
    )

    document.add_heading("2.2 技术栈与职责", level=2)
    add_table(
        document,
        ["技术", "在项目中的实际职责"],
        [
            ["FastAPI", "路由、请求校验、依赖注入、单端口静态页面服务"],
            ["Pydantic", "工单创建、更新、分析请求和响应模型"],
            ["SQLite", "父工单、子工单、快照持久化，WAL 并发读取"],
            ["LangGraph", "工单草稿分析节点编排与状态传递"],
            ["CopilotKit", "页面浮层、表单上下文、前端工具注册和运行时握手"],
            ["ag-ui-langgraph", "把 LangGraph 图暴露为 CopilotKit 可调用的 AG-UI 服务"],
            ["LlamaIndex Core", "使用 TextNode、NodeWithScore 和 QueryBundle 统一检索结构"],
            ["BM25S", "对 SOP 文档和历史工单进行本地相关性排序"],
            ["Tailwind CSS", "页面样式，已预构建为本地静态文件"],
            ["JWT", "REST 和 WebSocket 访问控制"],
        ],
        widths=[4.2, 11.3],
    )

    document.add_heading("3. 系统架构", level=1)
    add_image(document, ARCHITECTURE, 16.8, "图 3-1 系统架构与模块边界")
    document.add_heading("3.1 分层说明", level=2)
    add_table(
        document,
        ["层次", "主要内容", "关键文件"],
        [
            ["浏览器端", "表单、看板、查询、Agent 状态、日志、页面助手", "frontend/index.html"],
            ["服务层", "REST、WebSocket、JWT、限流、日志、静态资源", "backend/app/main.py"],
            ["编排层", "分类、拆单、线程分派、SLA、聚合", "backend/app/agents"],
            ["分析层", "事实提取、完整性检查、检索、草稿输出", "backend/app/assistant"],
            ["数据层", "SQLite、快照、可选 Redis 镜像", "backend/app/state/store.py"],
            ["知识层", "SOP 文档、历史工单、技能目录", "knowledge/、skills/"],
        ],
        widths=[3.2, 8.3, 5.0],
    )
    add_callout(
        document,
        "单端口设计的意义",
        "serve.py 同时挂载 FastAPI API 和 frontend 静态目录，启动后只需要一个端口。"
        "这样减少了跨域、代理和多进程启动问题，也让 Windows 一键脚本更可靠。",
        LIGHT_GREEN,
    )

    document.add_heading("4. 端到端工作流程", level=1)
    add_image(document, WORKFLOW, 16.8, "图 4-1 工单处理完整流程")
    document.add_heading("4.1 提交链路", level=2)
    add_table(
        document,
        ["步骤", "代码入口", "关键动作"],
        [
            ["1. 提交工单", "POST /api/v1/tickets/submit", "Pydantic 校验标题、描述、优先级"],
            ["2. 分类", "Orchestrator.handle_submission", "调用 classify_with_fallback"],
            ["3. 建父单", "TicketStore.save_ticket", "写入 tickets 表"],
            ["4. 拆子单", "_build_sub_tickets", "按部门去重并生成子工单"],
            ["5. 分派", "AgentRegistry.dispatch", "创建后台线程运行对应 Agent"],
            ["6. 返回", "SubmitResponse", "返回父单号、主部门、子工单号"],
        ],
        widths=[3.4, 5.8, 6.3],
    )
    document.add_heading("4.2 状态机", level=2)
    add_code(
        document,
        "created -> processing -> resolved\n"
        "                    -> waiting_info -> processing -> resolved\n"
        "                    -> escalated\n\n"
        "父工单：processing -> resolved（全部子单完成）\n"
        "父工单：processing -> escalated（存在升级/失败且全部子单终止）",
    )
    document.add_heading("4.3 为什么父工单不能由单个子任务提前结束", level=2)
    document.add_paragraph(
        "BaseAgent 在每次子任务完成后都会重新读取全部兄弟子工单。只有所有子工单都处于 resolved/closed 时，"
        "父工单才会变为 resolved；只要仍有 processing 或 waiting_info，父工单就保持 processing。"
        "这是多 Agent 场景中常见的状态一致性问题。"
    )

    document.add_heading("5. 后端模块组成", level=1)
    add_table(
        document,
        ["路径", "职责", "面试时可讲的点"],
        [
            ["app/main.py", "创建 FastAPI 应用、中间件、路由和 LangGraph endpoint", "应用工厂、生命周期、单端口集成"],
            ["app/api/tickets.py", "工单 CRUD、补充信息、聚合结果", "404、PATCH 校验、级联删除"],
            ["app/api/tickets_ws.py", "WebSocket 状态推送", "状态差异检测、终止条件和 JWT"],
            ["app/agents/orchestrator.py", "分类、拆单、SLA、聚合", "多 Agent 编排核心"],
            ["app/agents/base.py", "Agent 模板方法", "生命周期复用、异常升级、父状态收敛"],
            ["app/agents/registry.py", "Agent 注册和任务线程", "线程调度与 watchdog"],
            ["app/assistant/workflow.py", "LangGraph 草稿分析", "节点式工作流、上下文和工具调用"],
            ["app/assistant/retrieval.py", "SOP 与历史工单检索", "BM25、LlamaIndex 数据结构"],
            ["app/state/store.py", "SQLite 数据访问", "WAL、锁、乐观锁、快照"],
            ["app/middleware/__init__.py", "JWT、限流、日志", "请求链路治理"],
            ["app/core/llm.py", "分类回退、缺失字段检查、模型调用", "可用性与降级设计"],
        ],
        widths=[5.1, 6.6, 4.8],
    )

    document.add_heading("6. 多 Agent 设计", level=1)
    document.add_heading("6.1 Agent 分工", level=2)
    add_table(
        document,
        ["Agent", "二级意图", "典型输出"],
        [
            ["IT Agent", "账号、硬件、网络、软件", "权限申请、维修队列、网络诊断、软件安装"],
            ["HR Agent", "假期、考勤、证明、政策", "假期余额、考勤摘要、证明任务、HRBP 转交"],
            ["Procurement Agent", "办公用品、供应商、发票", "库存、审批层级、采购消息"],
            ["Finance Agent", "报销、预算、发票", "报销复核、预算余额、发票验真任务"],
            ["General Agent", "低置信度、无明确部门", "追问、人工队列、上下文快照"],
        ],
        widths=[4.0, 5.0, 7.5],
    )
    document.add_heading("6.2 编排器的职责边界", level=2)
    add_bullets(
        document,
        [
            "只负责识别意图、拆分任务、调度和聚合，不承担部门业务细节。",
            "每个部门 Agent 只处理自己的子工单，减少规则耦合。",
            "父工单保存全局状态，子工单保存部门、意图、结果和独立版本。",
            "Agent 异常会转换为 escalated，并保留错误信息供人工处理。",
        ],
    )
    document.add_heading("6.3 线程与 SLA", level=2)
    document.add_paragraph(
        "AgentRegistry 在应用 lifespan 启动时创建 watchdog 线程，每 5 秒调用一次 scan_sla_breaches。"
        "每次分派子工单时创建一个 daemon 线程，避免阻塞 FastAPI 工作线程。"
        "扫描器按部门 SLA 判断是否升级，waiting_info 使用额外宽限期，避免用户正在补充信息时被立即升级。"
    )
    add_callout(
        document,
        "面试可主动说明的局限",
        "当前线程模型适合本地单进程演示，但不具备跨进程消费、消息确认、重试队列和水平扩展能力。"
        "如果用于生产，应将 dispatch 和 SLA 扫描迁移到 RabbitMQ、Celery 或专用任务平台。",
        LIGHT_AMBER,
    )

    document.add_heading("7. LangGraph 工单分析工作流", level=1)
    add_table(
        document,
        ["节点", "输入", "主要逻辑", "输出"],
        [
            ["extract", "标题、描述、优先级", "分类、类型、标签、人员、截止时间", "draft / analysis / skills"],
            ["check_completeness", "draft", "按部门检查必填和建议字段", "missing_fields"],
            ["retrieve_context", "draft", "拼接查询并检索 SOP 与历史工单", "knowledge / similar_tickets"],
            ["compose", "前三个节点结果", "生成面向用户的草稿说明与追问", "assistant_message"],
        ],
        widths=[3.2, 4.2, 6.2, 3.0],
    )
    document.add_heading("7.1 事实约束", level=2)
    document.add_paragraph(
        "草稿节点只整理用户输入，不会伪造资产编号、预算、审批状态或人员。"
        "缺失字段通过 missing_fields 单独输出，前端在用户确认后继续补充。"
    )
    document.add_heading("7.2 CopilotKit 工具边界", level=2)
    document.add_paragraph(
        "页面助手注册的唯一写前端工具是 fillTicketForm。该工具派发 ticket-assistant:fill 事件，"
        "由表单页面更新可见字段。没有任何工具调用 submit、update 或 delete API。"
        "因此助手只能辅助编辑，不能绕过人工确认。"
    )
    add_callout(
        document,
        "为什么这个边界重要",
        "自然语言模型可能误判用户意图。把写操作限制在可编辑表单，而不是直接调用后端写接口，"
        "可以显著降低误提交和不可逆操作风险。",
        LIGHT_BLUE,
    )

    document.add_heading("8. Skills 的设计与真实作用", level=1)
    document.add_paragraph(
        "系统递归发现 skills 目录中的 SKILL.md，并读取 frontmatter 中的 name、description 和 version。"
        "目录中包含 15 个 Superpowers 技能和 1 个本地 to-tickets 兼容技能。"
    )
    add_table(
        document,
        ["技能来源", "数量", "项目中的使用方式"],
        [
            ["obra/superpowers", "15", "按固定提交 vendored，可在技能接口中查询元数据"],
            ["本地 to-tickets", "1", "提供会话转工单规则，核心规则在 workflow.py 中显式实现"],
        ],
        widths=[5.0, 2.0, 9.0],
    )
    document.add_heading("8.1 需要如实说明的边界", level=2)
    add_bullets(
        document,
        [
            "技能目录负责发现和展示，不会执行仓库中的任意脚本。",
            "工作流选择 to-tickets、brainstorming、systematic-debugging、verification-before-completion 作为阶段标记。",
            "工单字段抽取和缺失信息判断位于项目代码中，不能宣称全部由 Superpowers 自动完成。",
            "mattpocock/to-tickets 原始地址当时不可公开访问，仓库内是本地兼容实现，不是上游源码。",
        ],
    )

    document.add_heading("9. 数据模型与并发控制", level=1)
    add_table(
        document,
        ["表", "关键字段", "用途"],
        [
            ["tickets", "id、description、status、primary_dept、version、meta", "父工单和分类结果"],
            ["sub_tickets", "id、parent_id、dept、intent_label、status、result", "部门子工单"],
            ["snapshots", "ticket_id、seq、payload、captured_at", "变更快照和诊断"],
        ],
        widths=[3.5, 8.0, 5.0],
    )
    document.add_heading("9.1 SQLite WAL", level=2)
    document.add_paragraph(
        "SQLite 连接使用 WAL 日志模式和 NORMAL 同步级别，适合读多写少的本地应用。"
        "TicketStore 使用进程内锁串行化写操作，避免同一个连接被多线程同时写。"
    )
    document.add_heading("9.2 乐观锁", level=2)
    document.add_paragraph(
        "tickets 和 sub_tickets 都有 version 字段。更新父工单时可以选择传入 expected_version，"
        "SQL 会附加 version = ? 条件，rowcount 为 0 表示发生并发冲突或记录不存在。"
    )
    document.add_heading("9.3 删除一致性", level=2)
    document.add_paragraph(
        "删除父工单前先查询是否存在，随后在同一个事务内删除 snapshots、sub_tickets 和 tickets。"
        "任一步失败都会 rollback。删除不存在的工单返回 404，前端删除按钮也会二次确认。"
    )

    document.add_heading("10. API 与安全", level=1)
    add_table(
        document,
        ["安全点", "实现"],
        [
            ["HTTP 鉴权", "JWT Bearer，中间件统一解析 sub 和 roles"],
            ["WebSocket 鉴权", "token 查询参数校验，失败以 1008 关闭"],
            ["输入校验", "Pydantic 长度、枚举和自定义 meaningful 检查"],
            ["更新校验", "拒绝 null 和空字段，至少要求一个更新项"],
            ["删除校验", "记录不存在返回 404"],
            ["限流", "每客户端 20 个突发请求，按 5 rps 补充"],
            ["追踪", "X-Trace-Id 响应头和请求日志"],
            ["默认密钥", "只用于本地开发，部署通过 AGENT_JWT_SECRET 覆盖"],
        ],
        widths=[4.0, 11.5],
    )
    add_callout(
        document,
        "演示环境的诚实说明",
        "/api/v1/auth/token 允许直接签发 JWT，是为了本地演示和自动化测试。"
        "生产环境必须接入用户目录、OIDC 或统一身份平台，不能保留该公开签发入口。",
        LIGHT_AMBER,
    )

    document.add_heading("11. 前端组成", level=1)
    add_image(document, DESKTOP, 16.3, "图 11-1 桌面端工单看板")
    document.add_heading("11.1 页面板块", level=2)
    add_table(
        document,
        ["板块", "核心交互"],
        [
            ["提交工单", "新建、编辑、草稿预检查、提交"],
            ["工单看板", "状态统计、刷新、进入详情、编辑、删除"],
            ["进度查询", "查询父工单、父子状态、追加缺失信息"],
            ["Agent 状态", "查看在线状态、处理量和 uptime"],
            ["实时日志", "显示操作日志、WebSocket 事件、错误信息"],
            ["页面助手", "读取表单、对话、工具回填、人工复核"],
        ],
        widths=[4.0, 11.5],
    )
    document.add_heading("11.2 离线静态资源", level=2)
    document.add_paragraph(
        "前端不再依赖 Tailwind CDN。tailwind.config.js 和 tailwind.input.css 只用于开发时构建，"
        "运行时加载预构建的 tailwind.css、copilot-widget.css 和 copilot-widget.js。"
        "因此目标电脑不需要 Node.js，断网也能正常展示页面样式。"
    )
    add_image(document, CHAT, 16.3, "图 11-2 页面助手回填草稿")
    document.add_heading("11.3 移动端适配", level=2)
    document.add_paragraph(
        "导航在小屏切换为可横向滚动的标签条，统计卡片变为两列，表格保持横向滚动并固定操作列。"
        "页面助手宽度通过 min(420px, calc(100vw - 24px)) 限制，避免超出视口。"
    )
    add_image(document, MOBILE, 7.0, "图 11-3 移动端工单看板")

    document.add_heading("12. 关键工程决策", level=1)
    decisions = [
        (
            "为什么用规则回退",
            "外部模型可能超时、限流或没有配置 Key。关键词规则保证离线演示和测试可重复，"
            "模型只作为增强而不是单点依赖。",
        ),
        (
            "为什么用 BM25 而不是向量数据库",
            "项目规模小，SOP 文件数量有限。BM25 不需要额外服务，启动成本低，检索结果可解释。"
            "LlamaIndex 数据结构保留了后续替换向量检索的接口形态。",
        ),
        (
            "为什么把 WebSocket 做轮询式差异推送",
            "浏览器 WebSocket 连接只建立一次，服务端每秒比较子工单状态并只发送状态快照和迁移事件。"
            "实现简单，足够覆盖当前线程模型，也能明确终止条件。",
        ),
        (
            "为什么保留 SQLite",
            "一键部署和换电脑运行是当前优先级。SQLite 没有外部服务依赖，WAL 能满足本地并发读取。"
            "如果需要多实例，再迁移 PostgreSQL 或 MySQL 更合理。",
        ),
        (
            "为什么预构建前端",
            "目标电脑不需要 Node.js 和构建链，减少环境问题。开发源码仍保留，便于修改后重新构建。",
        ),
    ]
    for title, body in decisions:
        document.add_heading(title, level=2)
        document.add_paragraph(body)

    document.add_heading("13. 测试与验证", level=1)
    add_table(
        document,
        ["验证层", "内容", "结果"],
        [
            ["后端测试", "健康检查、CRUD、404、技能、AG-UI、WebSocket、SLA UTC", "13 passed"],
            ["静态检查", "Ruff F/I/UP/B 规则", "All checks passed"],
            ["依赖检查", "pip check 与应用模块导入", "通过"],
            ["浏览器 CRUD", "创建、编辑、删除、Dashboard 渲染", "通过"],
            ["页面助手", "info 200、run 200、fillTicketForm 回填", "通过"],
            ["响应式截图", "桌面与移动端", "通过"],
        ],
        widths=[3.2, 9.8, 2.5],
    )
    add_code(
        document,
        "cd backend\n"
        "$env:AGENT_LLM_ENABLED='0'\n"
        "$env:AGENT_SQLITE_PATH='data\\test_tickets.sqlite3'\n"
        ".\\.venv\\Scripts\\python.exe -m pytest tests -q\n\n"
        ".\\.venv\\Scripts\\ruff.exe check app tests",
    )

    document.add_heading("14. 五分钟演示脚本", level=1)
    steps = [
        "打开首页，介绍五个页面板块和右侧页面助手。",
        "输入包含 IT、HR、财务的混合请求，展示多意图拆分和父单号。",
        "进入进度查询，展示子工单、等待补充、实时状态和聚合结果。",
        "在工单看板点击编辑，修改标题和优先级，再点击删除并展示确认框。",
        "打开页面助手，输入一个硬件故障，展示助手只回填表单，不自动提交。",
        "打开 Swagger，展示 PATCH、DELETE、404 和 WebSocket 路由。",
        "最后展示 pytest、Ruff 和浏览器回归结果。",
    ]
    for index, step in enumerate(steps, start=1):
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{index}. ").bold = True
        paragraph.add_run(step)

    document.add_heading("15. 高频面试问题与回答", level=1)
    qa = [
        (
            "为什么使用多 Agent，而不是一个分类函数？",
            "分类函数只负责判断部门，多 Agent 还包含部门业务规则、结果结构、缺失信息处理、"
            "错误升级和状态回传。拆成独立 Agent 后，编排和业务细节解耦，新增部门时可以按统一接口扩展。",
        ),
        (
            "多意图是如何拆分的？",
            "classify_with_fallback 返回 sub_intents，每项包含 dept、confidence 和 span。"
            "Orchestrator 按 dept 去重，主部门保证存在，再为每个意图生成独立子工单和部门 SLA。",
        ),
        (
            "没有大模型 Key 时项目还能运行吗？",
            "可以。LLM 调用失败或未配置时自动使用关键词分类；完整性检查、SLA、检索和 CRUD 都是本地逻辑。"
            "这保证了演示可重复性和离线可用性。",
        ),
        (
            "waiting_info 有什么价值？",
            "它把缺失信息从异常变成正式状态。子工单不会带着未知资产号或预算被错误处理，"
            "用户可以只补充必要字段，然后重新进入处理队列。",
        ),
        (
            "SLA 扫描中你处理过什么细节问题？",
            "时间字符串是 UTC，早期使用 time.mktime 会按本地时区解析，在中国时区可能导致工单刚创建就被认为超时。"
            "现在使用带 timezone.utc 的 datetime 解析，并增加了测试。",
        ),
        (
            "如何保证父工单状态准确？",
            "单个子任务结束时不能直接把父工单设为 resolved。BaseAgent 会重新读取全部兄弟子任务，"
            "只有全部终态才收敛父状态。",
        ),
        (
            "WebSocket 为什么还传 Token？",
            "浏览器 WebSocket API 不能像 fetch 一样方便设置 Authorization 头，因此使用短期 JWT 查询参数。"
            "服务端验证失败会发送 error 并以 1008 关闭。",
        ),
        (
            "删除工单如何避免残留数据？",
            "TicketStore.delete_ticket 在一个事务内删除 snapshots、sub_tickets 和 tickets；"
            "记录不存在时不执行删除并返回 404。前端删除前还有确认框。",
        ),
        (
            "LangGraph 在项目里承担什么？",
            "承担草稿分析的节点编排：extract、check_completeness、retrieve_context、compose。"
            "每个节点只负责一步，便于测试和扩展。",
        ),
        (
            "LlamaIndex 和 BM25 分别做什么？",
            "LlamaIndex 提供 TextNode、NodeWithScore 和 QueryBundle 的统一结构；BM25S 实际执行本地相关性排序。"
            "这样可以不依赖向量数据库和外部服务。",
        ),
        (
            "Skills 是否真的会自动执行？",
            "不能夸大。当前实现安全地发现 SKILL.md 元数据，并在分析结果中标出相关技能。"
            "工单字段提取和缺失信息判断由项目代码显式完成，不会执行技能目录里的任意脚本。",
        ),
        (
            "这个项目最大的局限是什么？",
            "当前是单进程本地实现，Agent 调度使用线程，不是消息队列；业务系统数据是演示数据；"
            "认证入口也是演示接口。下一步应接入消息队列、真实身份系统和外部业务 API。",
        ),
        (
            "如果并发量增加，如何演进？",
            "先把任务分派改为 RabbitMQ/Celery，SLA 改为延迟消息或定时任务；"
            "再将 SQLite 换成 PostgreSQL/MySQL，WebSocket 通过 Redis Pub/Sub 做多实例广播。",
        ),
    ]
    for index, (question, answer) in enumerate(qa, start=1):
        document.add_heading(f"Q{index}. {question}", level=2)
        paragraph = document.add_paragraph(answer)
        paragraph.paragraph_format.space_after = Pt(5)

    document.add_heading("16. 一分钟项目介绍模板", level=1)
    add_callout(
        document,
        "建议表述",
        "我做过一个多智能体工单协同系统，后端使用 FastAPI，前端是原生页面加 CopilotKit 浮层。"
        "项目解决的核心问题是用户一句话包含多个部门诉求时，传统工单无法自动拆单和跟踪。"
        "我实现了父工单到子工单的拆分、IT/HR/采购/财务 Agent、waiting_info 补充信息、SLA 升级、"
        "WebSocket 状态推送和完整 CRUD。页面助手通过 LangGraph 生成可编辑草稿，只回填表单，不自动提交。"
        "项目还做了离线回退、13 个后端测试、Ruff 静态检查和 Playwright 浏览器回归。"
        "我也明确保留了边界：当前是本地单进程实现，生产化需要消息队列、真实业务 API 和统一身份系统。",
        LIGHT_GREEN,
    )

    document.add_heading("17. 后续演进路线", level=1)
    add_table(
        document,
        ["优先级", "改进项", "原因"],
        [
            ["P0", "接入真实 OIDC / 用户系统", "替换公开演示 Token 接口"],
            ["P0", "把 Agent 调度迁移到消息队列", "支持确认、重试、死信和水平扩展"],
            ["P1", "用 PostgreSQL/MySQL 替换 SQLite", "支持多实例和更高并发"],
            ["P1", "接入真实资产、HR、库存和财务 API", "把演示数据变成真实业务闭环"],
            ["P1", "增加角色权限与操作审计", "区分员工、主管、部门处理人和管理员"],
            ["P2", "向量检索与增量索引", "提升大规模知识库检索质量"],
            ["P2", "前端组件化与设计系统", "降低页面脚本维护成本"],
            ["P2", "CI 自动执行测试和构建", "保证 GitHub 提交质量"],
        ],
        widths=[2.3, 7.8, 5.4],
    )

    document.add_heading("18. 代码地图", level=1)
    add_code(
        document,
        "backend/app/main.py                     FastAPI 应用与路由装配\n"
        "backend/app/api/tickets.py              工单 CRUD 与补充信息\n"
        "backend/app/api/tickets_ws.py           WebSocket 状态推送\n"
        "backend/app/agents/orchestrator.py      分类、拆单、SLA、聚合\n"
        "backend/app/agents/base.py              Agent 生命周期模板\n"
        "backend/app/assistant/workflow.py       LangGraph 草稿分析\n"
        "backend/app/assistant/retrieval.py      SOP 与历史工单检索\n"
        "backend/app/state/store.py              SQLite 数据访问\n"
        "frontend/index.html                     业务页面与交互\n"
        "frontend/copilot-widget-src/src/index.jsx 页面助手与工具注册\n"
        "backend/tests/                          后端回归测试\n"
        "frontend/copilot-widget-src/verify-*.mjs 浏览器回归测试",
    )

    document.add_heading("19. 结论", level=1)
    document.add_paragraph(
        "这个项目的核心价值不在于堆叠框架，而在于把多部门工单拆单、状态流转、补充信息、"
        "SLA 升级、知识检索和人工确认组成一个能运行、能测试、能解释的完整闭环。"
        "面试讲解时应重点说明数据如何流转、为什么这样拆分模块、如何处理异常和边界，"
        "同时主动说明当前实现的生产化不足。"
    )
    add_callout(
        document,
        "面试原则",
        "只讲代码中真实存在并且能够演示的功能；主动区分已实现、演示数据和未来规划。"
        "能够准确说明边界，比把项目描述成生产系统更有说服力。",
        LIGHT_BLUE,
    )

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
